"""Load a policy document (PDF, DOCX, TXT, MD) into plain text with page markers."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Union

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


class PolicyLoadError(Exception):
    pass


def load_policy(source: Union[str, Path, bytes], filename: str | None = None) -> str:
    """Return the document text.

    `source` may be a path, or raw bytes (e.g. from a Streamlit upload) together
    with `filename` so the format can be detected.
    """
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise PolicyLoadError(f"File not found: {path}")
        data = path.read_bytes()
        filename = filename or path.name
    else:
        data = source

    if not filename:
        raise PolicyLoadError("A filename is needed to detect the document type.")

    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise PolicyLoadError(
            f"Unsupported file type '{ext}'. Use one of: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext == ".pdf":
        text = _load_pdf(data)
    elif ext == ".docx":
        text = _load_docx(data)
    else:
        text = _decode_text(data)

    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if len(text) < 50:
        raise PolicyLoadError(
            "The document has almost no readable text. If it is a scanned PDF, "
            "export it with OCR / as text first."
        )
    return text


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _load_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise PolicyLoadError("pypdf is not installed. Run: pip install -r requirements.txt") from exc

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            # Page markers let citations point to "page N" of the original PDF.
            pages.append(f"[[PAGE {number}]]\n{page_text}")
    return "\n\n".join(pages)


def _load_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise PolicyLoadError("python-docx is not installed. Run: pip install -r requirements.txt") from exc

    document = docx.Document(io.BytesIO(data))
    lines = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            lines.append("#" * min(int(level), 6) + " " + text)
        else:
            lines.append(text)
    # Tables (fee tables, baggage allowances) are common in policy docs.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)
